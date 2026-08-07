"""
sam_bulk — Bulk update item codes for widget items.

Supports:
  - upload: Update product codes for one or more widget items on Samaan
  - dry_run: Preview the CSV that would be uploaded (no actual upload)

CSV format for Samaan bulk upload:
  widget_item,item_code,priority
  slug_1,90513,1
  slug_1,90518,2
"""

import json
import os


async def handle_bulk(arguments: dict, configs: dict) -> dict:
    """Route bulk actions to handlers."""
    action = arguments.get("action")

    if not action:
        return {"error": "bulk requires 'action' parameter. Options: upload, dry_run, upload_from_sheet"}

    if action == "upload":
        return await _upload(arguments, configs)
    elif action == "dry_run":
        return await _dry_run(arguments, configs)
    elif action == "upload_from_sheet":
        return await _upload_from_sheet(arguments, configs)
    else:
        return {"error": f"bulk.{action} is not a valid action. Options: upload, dry_run, upload_from_sheet"}


def _build_csv(items: list) -> tuple:
    """Build bulk upload CSV from items list.

    Args:
        items: list of {"slug": "...", "products": "code1,code2,..."}

    Returns:
        (csv_string, total_rows, item_count)
    """
    rows = []
    for item in items:
        slug = item.get("slug", "")
        products = item.get("products", "")
        if not slug or not products:
            continue
        codes = [c.strip() for c in str(products).split(",") if c.strip()]
        for priority, code in enumerate(codes, 1):
            rows.append(f"{slug},{code},{priority}")

    # CORRECT column name: widget_item (NOT widget_item_slug_name)
    header = "widget_item,item_code,priority"
    csv_str = header + "\n" + "\n".join(rows)
    return csv_str, len(rows), len(items)


async def _upload(args: dict, configs: dict) -> dict:
    """Bulk upload item codes to Samaan widget items via urllib (CSRF-safe)."""
    env = args.get("env")
    if not env:
        return {"status": "missing_env", "message": "Which environment? PROD or UAT?", "options": ["PROD", "UAT"]}

    items = args.get("items")

    if not items or not isinstance(items, list):
        return {
            "status": "missing_fields",
            "message": "Provide 'items' — list of widget items with product codes.",
            "example": {
                "items": [
                    {"slug": "bau_plp_firstfold_sale_grocery_cl_wi_all_both", "products": "90513,90518,368"},
                ]
            }
        }

    confirm = args.get("confirm", False)

    csv_str, total_rows, item_count = _build_csv(items)

    if not confirm:
        preview_lines = csv_str.split("\n")[:20]
        return {
            "status": "ready",
            "message": f"Ready to upload {total_rows} product codes across {item_count} widget items on {env}.",
            "environment": env,
            "total_product_codes": total_rows,
            "widget_items": item_count,
            "csv_preview": "\n".join(preview_lines) + ("\n..." if len(csv_str.split("\n")) > 20 else ""),
            "next_step": "Confirm? Call again with confirm=true",
        }

    if env == "PROD" and not args.get("prod_ack"):
        return {
            "status": "prod_confirmation_required",
            "message": f"PROD BULK UPLOAD — {total_rows} codes across {item_count} items.\nCall again with prod_ack=true.",
            "environment": "PROD",
        }

    # Upload via urllib (not aiohttp — CSRF works with urllib)
    import urllib.request, urllib.parse, http.cookiejar, io, uuid

    samaan_cfg = configs.get("samaan", {})
    env_key = "PROD" if env == "PROD" else "UAT"
    base = samaan_cfg.get("environments", {}).get(env_key, "")
    username = os.environ.get(f"SAMAAN_{env_key}_USER") or samaan_cfg.get("credentials", {}).get(env_key, {}).get("username", "")
    password = os.environ.get(f"SAMAAN_{env_key}_PASS") or samaan_cfg.get("credentials", {}).get(env_key, {}).get("password", "")

    try:
        jar = http.cookiejar.CookieJar()
        opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
        opener.open(urllib.request.Request(f"{base}/login/", headers={'User-Agent': 'SAM-Bot/1.0'}), timeout=30)
        csrf = None
        for c in jar:
            if c.name == 'csrftoken': csrf = c.value
        opener.open(urllib.request.Request(f"{base}/login/",
            data=urllib.parse.urlencode({'csrfmiddlewaretoken': csrf, 'username': username, 'password': password}).encode(),
            headers={'Referer': f"{base}/login/", 'User-Agent': 'SAM-Bot/1.0'}), timeout=30)
        csrf = session = None
        for c in jar:
            if c.name == 'csrftoken': csrf = c.value
            if c.name == 'sessionid': session = c.value

        # Build multipart PUT
        boundary = uuid.uuid4().hex
        body = io.BytesIO()
        body.write(f"--{boundary}\r\nContent-Disposition: form-data; name=\"csrfmiddlewaretoken\"\r\n\r\n{csrf}\r\n".encode())
        body.write(f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"bulk.csv\"\r\nContent-Type: text/csv\r\n\r\n".encode())
        body.write(csv_str.encode())
        body.write(f"\r\n--{boundary}--\r\n".encode())

        url = f"{base}/api/app/bulk_upload_products_for_wi/"
        req = urllib.request.Request(url, data=body.getvalue(), method='PUT')
        req.add_header('Content-Type', f'multipart/form-data; boundary={boundary}')
        req.add_header('X-CSRFToken', csrf)
        req.add_header('Cookie', f'csrftoken={csrf}; sessionid={session}')
        req.add_header('Referer', f'{base}/widget-item/')
        req.add_header('User-Agent', 'SAM-Bot/1.0')

        resp = opener.open(req, timeout=120)
        result = json.loads(resp.read().decode())

        return {
            "status": "uploaded",
            "message": f"Bulk upload done on {env}! {total_rows} codes across {item_count} items.",
            "environment": env,
            "total_product_codes": total_rows,
            "widget_items": item_count,
            "samaan_response": result,
        }
    except urllib.error.HTTPError as e:
        return {"status": "failed", "error": f"HTTP {e.code}: {e.read().decode()[:200]}"}
    except Exception as e:
        return {"status": "failed", "error": str(e)}


async def _upload_from_sheet(args: dict, configs: dict) -> dict:
    """Bulk upload item codes from an Excel sheet with columns:
    Widget Item ID, Slug Name (widget_item_id), State, Sub Category, Rank, Item Code, Priority, ...

    Builds slug from state + sub_category: bau_HD_masthead_firstfold_june_2026_{state}_{SubCat}_sc_wi_all_both
    Or accepts custom slug_prefix to build: {slug_prefix}_{state}_{SubCat}_sc_wi_all_both

    Args:
        file_path: Path to .xlsx file
        sheet_name: Sheet name (default: first sheet)
        slug_prefix: Slug prefix (default: bau_HD_masthead_firstfold_june_2026)
        subcat_col: Column name for sub-category (default: Sub Category)
        state_col: Column name for state (default: State)
        item_code_col: Column name for item code (default: Item Code)
        rank_col: Column name for priority/rank (default: Rank)
        env: PROD or UAT
        confirm: true to upload
        prod_ack: true for PROD confirmation
    """
    import openpyxl

    env = args.get("env")
    if not env:
        return {"status": "missing_env", "message": "Which environment? PROD or UAT?"}

    file_path = args.get("file_path")
    if not file_path:
        return {"status": "missing_fields", "message": "Provide 'file_path' — path to .xlsx file"}

    try:
        wb = openpyxl.load_workbook(file_path)
    except Exception as e:
        return {"status": "failed", "error": f"Cannot open file: {e}"}

    sheet_name = args.get("sheet_name")
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.worksheets[0]

    slug_prefix = args.get("slug_prefix", "bau_HD_masthead_firstfold_june_2026")
    subcat_col = args.get("subcat_col", "Sub Category")
    state_col = args.get("state_col", "State")
    item_code_col = args.get("item_code_col", "Item Code")
    rank_col = args.get("rank_col", "Rank")

    # Find column indices from headers
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    col_subcat = headers.get(subcat_col)
    col_state = headers.get(state_col)
    col_code = headers.get(item_code_col)
    col_rank = headers.get(rank_col)

    if not all([col_subcat, col_state, col_code, col_rank]):
        return {
            "status": "missing_columns",
            "message": f"Required columns not found. Available: {list(headers.keys())}",
            "expected": [subcat_col, state_col, item_code_col, rank_col],
        }

    state_map = {'CG': 'cg', 'JH': 'jh', 'WB': 'wb'}
    csv_rows = []
    for r in range(2, ws.max_row + 1):
        subcat = ws.cell(r, col_subcat).value
        state = ws.cell(r, col_state).value
        item_code = ws.cell(r, col_code).value
        rank = ws.cell(r, col_rank).value
        if not subcat or not state or not item_code or not rank:
            continue
        st = state_map.get(str(state).upper(), str(state).lower())
        slug = f"{slug_prefix}_{st}_{subcat}_sc_wi_all_both"
        csv_rows.append((slug, int(item_code), int(rank)))

    csv_rows.sort(key=lambda x: (x[0], x[2]))
    csv_str = "widget_item,item_code,priority\n" + "\n".join(f"{s},{c},{p}" for s, c, p in csv_rows)
    slugs = set(r[0] for r in csv_rows)
    total_rows = len(csv_rows)
    item_count = len(slugs)

    if not args.get("confirm"):
        preview = csv_str.split("\n")[:20]
        return {
            "status": "ready",
            "message": f"Ready to upload {total_rows} product codes across {item_count} widget items on {env}.",
            "environment": env,
            "total_product_codes": total_rows,
            "widget_items": item_count,
            "sheet": ws.title,
            "csv_preview": "\n".join(preview) + ("\n..." if total_rows > 19 else ""),
            "next_step": "Call again with confirm=true",
        }

    if env == "PROD" and not args.get("prod_ack"):
        return {
            "status": "prod_confirmation_required",
            "message": f"PROD BULK UPLOAD — {total_rows} codes across {item_count} items.\nCall again with prod_ack=true.",
        }

    # Upload via urllib
    import urllib.request, urllib.parse, http.cookiejar, io, uuid

    samaan_cfg = configs.get("samaan", {})
    env_key = "PROD" if env == "PROD" else "UAT"
    base = samaan_cfg.get("environments", {}).get(env_key, "")
    username = os.environ.get(f"SAMAAN_{env_key}_USER") or samaan_cfg.get("credentials", {}).get(env_key, {}).get("username", "")
    password = os.environ.get(f"SAMAAN_{env_key}_PASS") or samaan_cfg.get("credentials", {}).get(env_key, {}).get("password", "")

    try:
        jar = http.cookiejar.CookieJar()
        opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
        opener.open(urllib.request.Request(f"{base}/login/", headers={'User-Agent': 'SAM-Bot/1.0'}), timeout=30)
        csrf = None
        for c in jar:
            if c.name == 'csrftoken': csrf = c.value
        opener.open(urllib.request.Request(f"{base}/login/",
            data=urllib.parse.urlencode({'csrfmiddlewaretoken': csrf, 'username': username, 'password': password}).encode(),
            headers={'Referer': f"{base}/login/", 'User-Agent': 'SAM-Bot/1.0'}), timeout=30)
        csrf = session = None
        for c in jar:
            if c.name == 'csrftoken': csrf = c.value
            if c.name == 'sessionid': session = c.value

        boundary = uuid.uuid4().hex
        body = io.BytesIO()
        body.write(f"--{boundary}\r\nContent-Disposition: form-data; name=\"csrfmiddlewaretoken\"\r\n\r\n{csrf}\r\n".encode())
        body.write(f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"bulk.csv\"\r\nContent-Type: text/csv\r\n\r\n".encode())
        body.write(csv_str.encode())
        body.write(f"\r\n--{boundary}--\r\n".encode())

        url = f"{base}/api/app/bulk_upload_products_for_wi/"
        req = urllib.request.Request(url, data=body.getvalue(), method='PUT')
        req.add_header('Content-Type', f'multipart/form-data; boundary={boundary}')
        req.add_header('X-CSRFToken', csrf)
        req.add_header('Cookie', f'csrftoken={csrf}; sessionid={session}')
        req.add_header('Referer', f'{base}/widget-item/')
        req.add_header('User-Agent', 'SAM-Bot/1.0')

        resp = opener.open(req, timeout=120)
        result = json.loads(resp.read().decode())

        return {
            "status": "uploaded",
            "message": f"Bulk upload from sheet done on {env}! {total_rows} codes across {item_count} items.",
            "environment": env,
            "total_product_codes": total_rows,
            "widget_items": item_count,
            "file": file_path,
            "sheet": ws.title,
            "samaan_response": result,
        }
    except urllib.error.HTTPError as e:
        return {"status": "failed", "error": f"HTTP {e.code}: {e.read().decode()[:200]}"}
    except Exception as e:
        return {"status": "failed", "error": str(e)}


async def _dry_run(args: dict, configs: dict) -> dict:
    """Preview bulk upload CSV without uploading."""
    env = args.get("env", "UAT")

    items = args.get("items")
    if not items or not isinstance(items, list):
        return {
            "status": "missing_fields",
            "message": "Provide 'items' — list of widget items with product codes.",
            "example": {
                "items": [
                    {"slug": "widget_item_slug", "products": "90513,90518"}
                ]
            }
        }

    csv_str, total_rows, item_count = _build_csv(items)

    return {
        "status": "preview",
        "message": f"Dry run — {total_rows} codes across {item_count} items. No upload.",
        "environment": env,
        "total_product_codes": total_rows,
        "widget_items": item_count,
        "csv": csv_str,
    }
